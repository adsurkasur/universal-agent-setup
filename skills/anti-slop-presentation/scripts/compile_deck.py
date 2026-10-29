import os
import sys
import argparse
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

def compile_html_to_pptx(html_path, output_pptx_path, slide_selector=".slide"):
    """
    Renders HTML slides styled with Tailwind/CSS into a high-resolution 16:9 PPTX deck.
    """
    if not os.path.exists(html_path):
        print(f"Error: HTML file not found at {html_path}")
        sys.exit(1)

    output_dir = os.path.dirname(os.path.abspath(output_pptx_path))
    os.makedirs(output_dir, exist_ok=True)
    temp_img_dir = os.path.join(output_dir, ".temp_slide_renders")
    os.makedirs(temp_img_dir, exist_ok=True)

    slide_images = []
    print(f"[1/3] Launching Headless Chromium to render slides...")
    
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1920, "height": 1080})
        normalized_url = f"file:///{os.path.abspath(html_path).replace(os.sep, '/')}"
        page.goto(normalized_url)
        # Wait for fonts & layout
        page.wait_for_timeout(2000)

        slides = page.locator(slide_selector).all()
        count = len(slides)
        print(f"[2/3] Found {count} slide(s) to capture...")

        for idx, slide_loc in enumerate(slides, start=1):
            img_path = os.path.join(temp_img_dir, f"slide_{idx:02d}.png")
            slide_loc.screenshot(path=img_path)
            slide_images.append(img_path)
            print(f"   -> Rendered Slide {idx}/{count} (1920x1080)")

        browser.close()

    print(f"[3/3] Compiling into 16:9 Widescreen PPTX...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    for img_path in slide_images:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, 0, 0, Inches(13.333), Inches(7.5))

    prs.save(output_pptx_path)
    print(f"[OK] Success! Presentation ready at: {output_pptx_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Compile HTML slides to 16:9 PPTX")
    parser.add_argument("html", help="Path to input HTML slide file")
    parser.add_argument("output", help="Path to output PPTX file")
    parser.add_argument("--selector", default=".slide", help="CSS selector for slides (default: .slide)")
    args = parser.parse_args()

    compile_html_to_pptx(args.html, args.output, args.selector)
